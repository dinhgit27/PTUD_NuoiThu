# -*- coding: utf-8 -*-
"""
Script chuẩn hóa và hoàn thiện file 'test-plan-iterations.xlsx' theo đúng chuẩn mực của cô giáo:
- Giữ nguyên cấu trúc 10 Sheets chuẩn:
  1. ReadMe
  2. Test Case List
  3. Test Case Coverage Plan
  4. Test Iteration 1
  5. Test Iteration 2
  6. Test Iteration 3
  7. Test Iteration 4
  8. Test Iteration 5
  9. Stats
  10. Pivot Tables
- Nạp chuẩn xác 25 kịch bản bao phủ toàn diện 10 phân hệ của OfficePet.
- Cột Classification A: Critical / Major / Minor / Trivial (Đúng ghi chú ô O4 của cô: 'cột A-B: major, minor, critical').
- Cột Classification B: P1 - High / P2 - Medium / P3 - Low.
- Font: 100% Times New Roman cho toàn bộ các ô trong file.
- Không che chữ: Tự động căn chỉnh độ rộng cột và độ cao hàng có wrap_text.
- Không tạo dữ liệu ảo: Trạng thái trong Iterations là 'Not Run' (hoặc 'N/A' nếu không nằm trong đợt), các cột Defect/Comment để trống.
"""

import sys
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEMPLATE_PATH = r"c:\Users\PHI HUNG\Downloads\PTUD_NuoiThu\testplan\test-plan-iterations.xlsx"

# 25 Kịch bản kiểm thử chuẩn hóa bao phủ 10 phân hệ của OfficePet
OFFICEPET_25_TEST_CASES = [
    # --- PHÂN HỆ 01: XÁC THỰC & QUẢN LÝ TÀI KHOẢN (TS01) ---
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.1",
        "Đăng ký tài khoản mới hợp lệ & Kiểm thực lỗi email/mật khẩu",
        "[Mục đích] Xác nhận đăng ký tài khoản nhân viên.\n[Điều kiện] Chưa có tài khoản, ở trang Đăng ký.\n[Các bước] 1. Nhập Họ tên 'Nguyễn Văn A', Email hợp lệ, Mật khẩu >= 8 ký tự, chọn Phòng ban. 2. Nhấn 'Đăng ký'. 3. Thử lại với email trùng hoặc pass < 8 ký tự.\n[Kỳ vọng] Tạo tài khoản mã hóa bcrypt, báo thành công; chặn email trùng và mật khẩu yếu.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-01",
        "FR-01.1, FR-01.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.2",
        "Đăng nhập phân quyền vai trò (Role ADMIN & USER) và JWT Token",
        "[Mục đích] Xác thực phân quyền truy cập giữa Quản trị viên và Nhân viên.\n[Điều kiện] Có tài khoản Admin và User trong hệ thống.\n[Các bước] 1. Đăng nhập User -> vào Khu vườn thú cưng. 2. Đăng xuất. 3. Đăng nhập Admin -> hiển thị menu Admin Console.\n[Kỳ vọng] JWT token cấp đúng role; User bị chặn vào /admin, Admin vào được trang quản trị.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-01",
        "FR-01.3, FR-01.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.3",
        "Quên mật khẩu qua mã OTP Email & Cập nhật ảnh đại diện Avatar",
        "[Mục đích] Kiểm tra khôi phục mật khẩu và đổi thông tin cá nhân.\n[Điều kiện] Tài khoản có email hợp lệ.\n[Các bước] 1. Bấm 'Quên mật khẩu', nhập email nhận OTP 6 số. 2. Nhập OTP và đổi mật khẩu mới. 3. Đăng nhập đổi avatar trong trang Cá nhân.\n[Kỳ vọng] OTP gửi thành công, đổi mật khẩu và đổi avatar cập nhật tức thì vào DB.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-01",
        "FR-01.5, FR-01.6",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),

    # --- PHÂN HỆ 02: NUÔI & TƯƠNG TÁC THÚ CƯNG 11 HOẠT ẢNH (TS02) ---
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.1",
        "Khởi tạo nhận nuôi thú cưng & Cơ chế suy giảm chỉ số sinh học tự nhiên",
        "[Mục đích] Xác thực nhận nuôi thú ban đầu và suy giảm chỉ số Đói/Khát/Năng lượng theo thời gian.\n[Điều kiện] Tài khoản mới đăng nhập lần đầu.\n[Các bước] 1. Chọn thú cưng mặc định (Corgi), đặt tên 'Bắp Cải', nhận nuôi. 2. Quan sát chỉ số Đói/Khát giảm dần sau mỗi 5-10 phút.\n[Kỳ vọng] UserPet được tạo với Level 1, 100 Năng lượng, 80 Đói/Khát; các thanh chỉ số đổi màu cảnh báo khi < 30%.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-02",
        "FR-02.1, FR-02.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.2",
        "Hành động Cho ăn (EATING) và Cho uống nước (DRINKING) tách bạch",
        "[Mục đích] Kiểm tra hoạt ảnh chuẩn hóa riêng biệt giữa Ăn và Uống theo chuẩn 11 hoạt ảnh mới.\n[Điều kiện] Thú cưng có Đói < 80 và Khát < 80.\n[Các bước] 1. Nhấn nút 'Cho ăn' -> quan sát hoạt ảnh EATING. 2. Nhấn nút 'Cho uống nước' -> quan sát hoạt ảnh DRINKING riêng biệt.\n[Kỳ vọng] Cho ăn hồi +30 Đói với hoạt ảnh nhai thức ăn; Cho uống hồi +25 Khát với hoạt ảnh uống nước riêng; kết thúc trở về IDLE.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-02",
        "FR-02.3, FR-02.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.3",
        "Tương tác vuốt ve (PETTING) giới hạn +30 EXP/ngày và Hoạt ảnh HAPPY / SAD",
        "[Mục đích] Kiểm tra tương tác vuốt ve, giới hạn chống cày cấp và chuyển đổi hoạt ảnh cảm xúc.\n[Điều kiện] Thú cưng đang ở trạng thái IDLE.\n[Các bước] 1. Click chuột vuốt ve thú cưng liên tục 6 lần (+5 EXP/lần). 2. Click tiếp tục lần 7 trở đi. 3. Để đói khát < 15%.\n[Kỳ vọng] Thú cưng hiển thị tim (PETTING), cộng tối đa +30 EXP/ngày; click tiếp vẫn vui nhưng không cộng thêm EXP; khi đói hiển thị HUNGRY/SAD.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-02",
        "FR-02.5, FR-02.6",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.4",
        "Cho thú cưng ngủ (SLEEPING) hồi năng lượng và Thăng cấp thú cưng",
        "[Mục đích] Kiểm tra cơ chế hồi phục năng lượng và thăng cấp level.\n[Điều kiện] Năng lượng < 50, EXP đạt 95/100.\n[Các bước] 1. Bấm 'Đi ngủ' -> phòng tối, hoạt ảnh SLEEPING phập phồng chữ Zzz. 2. Đánh thức. 3. Thực hiện chăm sóc để đạt 100 EXP.\n[Kỳ vọng] Năng lượng tăng dần khi ngủ; khi đạt 100 EXP kích hoạt Level Up, hoạt ảnh nhảy múa DANCING và nhận thưởng.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-02",
        "FR-02.7, FR-02.8",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),

    # --- PHÂN HỆ 03: DESKTOP COMPANION & WIDGET PIP (TS03) ---
    (
        "TS03",
        "Desktop Companion & PiP",
        "TC3.1",
        "Thu nhỏ thú cưng ra màn hình Desktop (PiP) và Kéo thả ghi nhớ tọa độ",
        "[Mục đích] Kiểm tra tính năng Document PiP luôn nổi trên desktop và ghi nhớ vị trí.\n[Điều kiện] Trình duyệt hỗ trợ Document PiP API.\n[Các bước] 1. Nhấn nút Desktop PiP trên web. 2. Kéo thả thú cưng sang góc phải màn hình. 3. Đóng và mở lại PiP.\n[Kỳ vọng] Cửa sổ popup trong suốt nổi trên màn hình; khi mở lại tự động khôi phục đúng tọa độ đã kéo thả.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-03",
        "FR-03.1, FR-03.2",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS03",
        "Desktop Companion & PiP",
        "TC3.2",
        "Thú cưng tự do đi lại trong ranh giới (Roaming) và Menu thao tác nhanh",
        "[Mục đích] Kiểm tra chuyển động đi bộ tự nhiên và bảng điều khiển chuột phải.\n[Điều kiện] Cửa sổ Desktop đang mở.\n[Các bước] 1. Bật chế độ Roaming, quan sát thú đi dạo. 2. Nhấp chuột phải lên thú cưng. 3. Nhấn 'Cho ăn nhanh'.\n[Kỳ vọng] Thú cưng quay đầu khi chạm mép màn hình (WALKING); menu chuột phải mở mượt mà và thực hiện hành động tức thì.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-03",
        "FR-03.3, FR-03.4",
        "Minor",
        "P3 - Low",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),

    # --- PHÂN HỆ 04: ĐỒNG HÀNH NĂNG SUẤT POMODORO & WELLNESS (TS04) ---
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.1",
        "Phiên Pomodoro 25 phút, thú cưng gõ phím TYPING và Nhận thưởng hoàn thành",
        "[Mục đích] Kiểm tra bộ đếm giờ Pomodoro, đồng bộ trạng thái gõ phím và cộng điểm thưởng.\n[Điều kiện] Thú cưng ở trạng thái IDLE.\n[Các bước] 1. Bắt đầu phiên 25 phút -> thú chuyển sang TYPING. 2. Tạm dừng -> trở về IDLE. 3. Để đếm hết 25 phút.\n[Kỳ vọng] Đồng hồ chạy chuẩn xác; thú cưng ngồi gõ máy tính chăm chỉ; kết thúc phiên nhận +50 EXP, +20 Vàng và phát nhạc chuông.",
        "Approved",
        "Tester 3",
        "Lead QA",
        "",
        "UC-04",
        "FR-04.1, FR-04.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 3"
    ),
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.2",
        "Chế độ Nghiêm khắc cảnh báo xao nhãng khi người dùng rời tab làm việc",
        "[Mục đích] Kiểm tra tính năng chống xao nhãng của Pomodoro.\n[Điều kiện] Bật Strict Mode trong cài đặt Pomodoro.\n[Các bước] 1. Bắt đầu phiên làm việc. 2. Chuyển sang tab khác (Youtube/MXH) hơn 15 giây. 3. Quay lại ứng dụng.\n[Kỳ vọng] Trình duyệt hiển thị thông báo nhắc nhở tập trung; thú cưng chuyển vẻ mặt buồn (SAD) cảnh báo.",
        "Approved",
        "Tester 3",
        "Lead QA",
        "",
        "UC-04",
        "FR-04.3",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 3"
    ),
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.3",
        "Thực hiện Bài tập thở thư giãn (4-7-8 / Box) và Trình phát nhạc Lofi",
        "[Mục đích] Kiểm tra các tính năng chăm sóc sức khỏe tinh thần và phát nhạc nền.\n[Điều kiện] Thiết bị có loa / tai nghe.\n[Các bước] 1. Chọn bài tập thở 4-7-8, làm theo nhịp thở. 2. Bật nhạc Lofi, điều chỉnh âm lượng, chuyển track tiếng mưa.\n[Kỳ vọng] Vòng tròn thở co giãn đúng nhịp giây; nhạc Lofi phát êm ái, chuyển bài mượt mà không lỗi âm thanh.",
        "Approved",
        "Tester 3",
        "Lead QA",
        "",
        "UC-04",
        "FR-04.4, FR-04.5",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 3"
    ),

    # --- PHÂN HỆ 05: CHATBOT AI TRÒ CHUYỆN ĐỒNG HÀNH (TS05) ---
    (
        "TS05",
        "Chatbot AI Đồng hành",
        "TC5.1",
        "Hộp thoại Chatbot AI, gửi câu hỏi công việc và Nhận phản hồi thông minh",
        "[Mục đích] Kiểm tra giao diện chat drawer và phản hồi từ AI Gemini.\n[Điều kiện] Kết nối Internet hoạt động, API Key được cấu hình.\n[Các bước] 1. Mở ngăn kéo chat bên cạnh thú cưng. 2. Nhập câu hỏi sắp xếp công việc. 3. Nhấn gửi.\n[Kỳ vọng] Tin nhắn hiển thị đẹp mắt, thú cưng hiển thị trạng thái suy nghĩ và trả lời động viên sâu sắc trong 1-3s.",
        "Approved",
        "Tester 3",
        "Lead QA",
        "",
        "UC-05",
        "FR-05.1, FR-05.2",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        "Tester 3"
    ),
    (
        "TS05",
        "Chatbot AI Đồng hành",
        "TC5.2",
        "Đồng bộ hoạt ảnh biểu cảm thú cưng (HAPPY / SAD) theo sắc thái trò chuyện",
        "[Mục đích] Kiểm tra sự gắn kết cảm xúc giữa Chat AI và hoạt ảnh thú cưng.\n[Điều kiện] Chatbot AI đang hoạt động.\n[Các bước] 1. Chat câu khoe tin vui thành tích -> quan sát thú. 2. Chat câu tâm sự buồn -> quan sát thú.\n[Kỳ vọng] Với tin vui thú cưng kích hoạt HAPPY / DANCING; với chuyện buồn thú làm động tác chia sẻ an ủi.",
        "Approved",
        "Tester 3",
        "Lead QA",
        "",
        "UC-05",
        "FR-05.3",
        "Minor",
        "P3 - Low",
        ("No", "No", "Yes", "Yes", "Yes"),
        "Tester 3"
    ),

    # --- PHÂN HỆ 06: NHIỆM VỤ & THI ĐUA PHÒNG BAN (TS06) ---
    (
        "TS06",
        "Nhiệm vụ & Thi đua",
        "TC6.1",
        "Nhiệm vụ hàng ngày, tự động cộng dồn tiến độ và Nhận thưởng EXP",
        "[Mục đích] Kiểm tra hệ thống tự động ghi nhận hành vi người dùng vào tiến độ nhiệm vụ.\n[Điều kiện] Nhiệm vụ 'Cho ăn 3 lần' có tiến độ 0/3.\n[Các bước] 1. Thực hiện cho thú ăn 3 lần. 2. Mở tab Nhiệm vụ. 3. Bấm 'Nhận thưởng'.\n[Kỳ vọng] Tiến trình tự động nhảy 3/3; nút 'Nhận thưởng' sáng lên; bấm nhận cộng đúng EXP và chuyển sang Đã nhận.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-06",
        "FR-06.1, FR-06.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS06",
        "Nhiệm vụ & Thi đua",
        "TC6.2",
        "Bảng xếp hạng thi đua cá nhân, phòng ban và Cơ chế reset tự động lúc 00:00",
        "[Mục đích] Kiểm tra xếp hạng thi đua và cơ chế reset ngày mới.\n[Điều kiện] Có dữ liệu điểm của các phòng ban.\n[Các bước] 1. Mở Bảng xếp hạng, lọc theo Tuần/Tháng/Phòng ban. 2. Kiểm tra reset nhiệm vụ sau 00:00.\n[Kỳ vọng] Xếp hạng sắp xếp đúng theo điểm số giảm dần; nhiệm vụ ngày cũ được làm mới lúc 00:00.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-06",
        "FR-06.3, FR-06.4",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),

    # --- PHÂN HỆ 07: KHU VƯỜN CỘNG ĐỒNG & HỆ THỐNG BẠN BÈ (TS07) ---
    (
        "TS07",
        "Khu vườn & Bạn bè",
        "TC7.1",
        "Khu vườn chung đa thú cưng phòng ban, Kết nối bạn bè & Ghé thăm vườn",
        "[Mục đích] Kiểm tra render nhiều thú cưng trên Canvas chung, tính năng kết bạn và thăm vườn.\n[Điều kiện] Đồng nghiệp trong phòng ban cùng online.\n[Các bước] 1. Mở 'Vườn Chung' xem thú cưng các đồng nghiệp. 2. Tìm kiếm email gửi kết bạn. 3. Ghé thăm vườn riêng của bạn bè.\n[Kỳ vọng] Thú cưng hiển thị đồng thời mượt mà không đè tọa độ; kết bạn thành công và chuyển cảnh thăm vườn chuẩn xác.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-07",
        "FR-07.1, FR-07.2",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),

    # --- PHÂN HỆ 08: CỬA HÀNG SHOP & KHO ĐỒ NỘI BỘ (TS08) ---
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.1",
        "Mua Thú cưng mới từ Shop bằng điểm EXP và Đổi thú cưng đại diện trong Kho đồ",
        "[Mục đích] Kiểm tra giao dịch mua thú cưng bằng tiền tệ tích lũy và chuyển đổi thú cưng hoạt động.\n[Điều kiện] Tài khoản có đủ 500 EXP.\n[Các bước] 1. Vào Shop tab Thú cưng, chọn 'Cáo Tuyết' (500 EXP), nhấn Mua. 2. Vào Kho đồ, bấm 'Đặt làm đại diện'.\n[Kỳ vọng] EXP bị trừ đúng 500; Cáo Tuyết được thêm vào UserPet; Màn hình chính và Desktop đổi ngay sang Cáo Tuyết.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-08",
        "FR-08.1, FR-08.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.2",
        "Mua Skin/Trang phục và Trang bị (Equip/Unequip) lên thú cưng",
        "[Mục đích] Kiểm tra tính năng mua sắm trang phục và hiển thị skin trên canvas.\n[Điều kiện] Người dùng có đủ Vàng và sở hữu thú cưng tương thích.\n[Các bước] 1. Mua skin 'Áo Hoodie Dev'. 2. Vào Kho đồ nhấn 'Trang bị' (Equip). 3. Bấm 'Tháo ra' (Unequip).\n[Kỳ vọng] Khi Equip thú cưng hiển thị áo hoodie mượt mà; khi Unequip trở về diện mạo mặc định.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-08",
        "FR-08.3, FR-08.4",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.3",
        "Mua bối cảnh Nền vườn mới và Áp dụng thay đổi không gian làm việc",
        "[Mục đích] Kiểm tra tùy biến bối cảnh nền và hiệu ứng gia tăng EXP nội tại.\n[Điều kiện] Đủ EXP mua background mới.\n[Các bước] 1. Mua nền 'Văn phòng Cyberpunk'. 2. Vào Kho đồ bấm 'Áp dụng làm nền vườn'.\n[Kỳ vọng] Khu vườn đổi ngay sang phong cách Cyberpunk, thú cưng đứng đúng vị trí mặt sàn; buff EXP hoạt động chính xác.",
        "Approved",
        "Tester 2",
        "Lead QA",
        "",
        "UC-08",
        "FR-08.5",
        "Minor",
        "P3 - Low",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        "Tester 2"
    ),

    # --- PHÂN HỆ 09: CỔNG QUẢN TRỊ ADMIN CONSOLE & FORGE TOOLS (TS09) ---
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.1",
        "Bảng điều khiển số liệu hệ thống và Đổi quyền người dùng (Role USER <-> ADMIN)",
        "[Mục đích] Kiểm tra các widget thống kê quản trị và phân quyền bảo mật.\n[Điều kiện] Đăng nhập tài khoản Admin.\n[Các bước] 1. Mở Admin Dashboard -> xem số liệu người dùng, thú cưng. 2. Vào Quản lý Người dùng, đổi role một tài khoản.\n[Kỳ vọng] Số liệu tải chính xác từ DB; role được cập nhật ngay lập tức; chặn tài khoản bị ban đăng nhập.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-09",
        "FR-09.1, FR-09.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.2",
        "Xưởng Pet Forge tạo thú 11 hành động, bỏ ô mô tả thừa, dọn sạch form và Hiện ở Shop",
        "[Mục đích] Xác thực cải tiến Pet Forge: preview đủ 11 hành động, không có ô mô tả thừa, dọn sạch UI sau tạo và lưu DB hiển thị ngay tại Shop.\n[Điều kiện] Admin đang ở công cụ Pet Forge.\n[Các bước] 1. Upload Sprite sheet 11 hàng. 2. Xác nhận form không có ô 'Mô tả'. 3. Nhập Tên, Giá và nhấn 'Tạo thú cưng'. 4. Kiểm tra form reset sạch sẽ và sang Quản lý Shop kiểm tra.\n[Kỳ vọng] Form dọn sạch giao diện ngay sau khi tạo; thú mới lưu vào bảng PetType và xuất hiện ngay lập tức trong Quản lý Shop.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-09",
        "FR-09.3, FR-09.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.3",
        "Xưởng tạo Skin, Xưởng tạo Vật phẩm hoạt ảnh và Chức năng Xóa thú cưng khỏi hệ thống",
        "[Mục đích] Xác thực việc tách riêng 2 xưởng tạo Skin/Vật phẩm, loại bỏ mục vật phẩm tĩnh cũ và kiểm tra nút xóa thú.\n[Điều kiện] Admin ở trang Quản lý Shop.\n[Các bước] 1. Kiểm tra menu điều hướng sang Tạo Skin và Tạo Vật phẩm (đã bỏ nút thêm vật phẩm cũ và mục vật phẩm tĩnh). 2. Nhấn nút Xóa một thú cưng thử nghiệm.\n[Kỳ vọng] Modal xác nhận xóa hiển thị an toàn; sau khi bấm xóa, thú cưng bị gỡ khỏi danh sách và database.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-09",
        "FR-09.5, FR-09.6",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),

    # --- PHÂN HỆ 10: PHI CHỨC NĂNG, BẢO MẬT & TƯƠNG THÍCH NGƯỢC (TS10) ---
    (
        "TS10",
        "Phi Chức Năng & Bảo Mật",
        "TC10.1",
        "Tương thích ngược thú cũ 17 hàng và Fallback CELEBRATING->HAPPY, EXERCISING->DANCING",
        "[Mục đích] Đảm bảo các con thú tạo trước đây (chuẩn 17 hàng) và thú mới (11 hàng) đều hiển thị bình thường không bị lỗi.\n[Điều kiện] Hệ thống có thú cưng sprite sheet 17 hàng cũ.\n[Các bước] 1. Mở xem thú cưng cũ 17 hàng. 2. Kích hoạt các sự kiện cũ CELEBRATING và EXERCISING.\n[Kỳ vọng] Thú cũ chạy mượt mà không lệch tọa độ frame; cơ chế fallback tự động chuyển CELEBRATING->HAPPY và EXERCISING->DANCING không gây crash console.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-10",
        "FR-10.1, FR-10.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
    (
        "TS10",
        "Phi Chức Năng & Bảo Mật",
        "TC10.2",
        "Render Canvas 60 FPS không rò rỉ RAM, Chặn Route Guard /admin và Chống XSS",
        "[Mục đích] Kiểm tra hiệu năng mượt mà của Canvas và kiểm thử an toàn bảo mật hệ thống.\n[Điều kiện] Đăng nhập bằng tài khoản User thông thường.\n[Các bước] 1. Đo FPS khi thú cưng chuyển động trong 30 phút. 2. Nhập URL /admin/pets trên trình duyệt. 3. Thử inject payload <script>alert(1)</script> vào tên thú.\n[Kỳ vọng] FPS đạt 55-60 FPS ổn định, RAM không phình; route /admin bị chặn 403; payload XSS bị escape an toàn.",
        "Approved",
        "Tester 1",
        "Lead QA",
        "",
        "UC-10",
        "FR-10.3, FR-10.4, FR-10.5",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        "Tester 1"
    ),
]

assert len(OFFICEPET_25_TEST_CASES) == 25, f"Số lượng phải đúng chuẩn 25, hiện tại là {len(OFFICEPET_25_TEST_CASES)}"

print(f"Tổng số test case chuẩn hóa: {len(OFFICEPET_25_TEST_CASES)} kịch bản.")

print(f"Đang mở workbook: {TEMPLATE_PATH}...")
wb = openpyxl.load_workbook(TEMPLATE_PATH)

FONT_TNR = "Times New Roman"

def apply_tnr_font(cell, base_font=None, size=None, bold=None, italic=None, color=None):
    if base_font:
        cell.font = Font(
            name=FONT_TNR,
            size=size if size is not None else base_font.size,
            bold=bold if bold is not None else base_font.bold,
            italic=italic if italic is not None else base_font.italic,
            color=color if color is not None else base_font.color
        )
    else:
        cell.font = Font(name=FONT_TNR, size=size or 11, bold=bool(bold), italic=bool(italic))

# =============================================================================
# 1. CẬP NHẬT SHEET 'ReadMe'
# =============================================================================
print("1. Đang cập nhật Sheet 'ReadMe'...")
ws_readme = wb['ReadMe']
ws_readme['B1'] = "Phi Hùng & Nhóm phát triển PTUD_NuoiThu"
ws_readme['B3'] = "Release v1.0.0 (Chuẩn hóa 11 Hoạt ảnh)"
ws_readme['B4'] = "24/10/2026"
ws_readme['B5'] = "OfficePet System - PTUD_NuoiThu © 2026"

readme_notes = [
    (7, "A", "Tên Đồ Án:"),
    (7, "B", "HỆ THỐNG NUÔI THÚ CƯNG ẢO ĐỒNG HÀNH NĂNG SUẤT VĂN PHÒNG (OFFICEPET)"),
    (8, "A", "Quy Chuẩn:"),
    (8, "B", "Kế hoạch kiểm thử lặp (Test Plan Iterations) theo đúng biểu mẫu phân kỳ của Giảng viên hướng dẫn."),
    (9, "A", "Phạm Vi:"),
    (9, "B", "Bao phủ đầy đủ 10 Phân hệ chính (TS01 -> TS10) với 25 kịch bản tích hợp và 5 Vòng lặp kiểm thử (Iterations)."),
    (11, "A", "Quy Ước Iteration:"),
    (11, "B", "Iteration 1: Smoke Test các luồng Critical\nIteration 2: Kiểm thử Core Features (Auth, Nuôi thú, Pomodoro, Admin Forge)\nIteration 3: Kiểm thử toàn diện chức năng (Full Functional: Shop, Vườn chung, Bạn bè, AI Chat)\nIteration 4: Kiểm thử Hồi quy & Phi chức năng (Regression, Tương thích thú cũ 17 hàng, Canvas 60fps, Bảo mật)\nIteration 5: Nghiệm thu toàn diện trước bàn giao (Final Acceptance / UAT)"),
    (17, "A", "Legend / Quy ước Trạng thái:"),
    (18, "B", "Not Run: Sẵn sàng thực thi, chưa chạy test (Tuyệt đối không tạo kết quả ảo khi chưa bắt đầu test)."),
    (19, "B", "Run - Passed: Đã chạy test thực tế và đạt kết quả mong đợi."),
    (20, "B", "Run - Failed: Đã chạy test thực tế và phát hiện lỗi lệch yêu cầu."),
    (21, "B", "Blocked: Bị chặn không thể test do lỗi phụ thuộc ở bước trước."),
    (22, "B", "N/A: Test case không nằm trong phạm vi của đợt Iteration này.")
]

for r, c, val in readme_notes:
    ws_readme[f"{c}{r}"] = val

# =============================================================================
# 2. CẬP NHẬT SHEET 'Test Case List'
# =============================================================================
print("2. Đang cập nhật Sheet 'Test Case List'...")
ws_tc = wb['Test Case List']

# Cập nhật danh sách lookup data validation chuẩn ở phía dưới
statuses_list = ['Not Started', 'In Progress', 'Draft', 'Pending Approval', 'Approved']
for i, s in enumerate(statuses_list):
    ws_tc.cell(row=31 + i, column=1, value=s)

class_a_list = ['Critical', 'Major', 'Minor', 'Trivial']
for i, ca in enumerate(class_a_list):
    ws_tc.cell(row=31 + i, column=4, value=ca)

class_b_list = ['P1 - High', 'P2 - Medium', 'P3 - Low', 'P4 - Lowest']
for i, cb in enumerate(class_b_list):
    ws_tc.cell(row=46 + i, column=4, value=cb)

owners_list = ['Tester 1', 'Tester 2', 'Tester 3', 'Lead QA']
for i, ow in enumerate(owners_list):
    ws_tc.cell(row=42 + i, column=1, value=ow)

# Nạp đúng 25 test cases vào hàng 2 -> 26
for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
    row_num = 2 + idx
    (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

    ws_tc.cell(row_num, 1, sc_id)
    ws_tc.cell(row_num, 2, sc_name)
    ws_tc.cell(row_num, 3, tc_id)
    ws_tc.cell(row_num, 4, tc_name)
    ws_tc.cell(row_num, 5, tc_desc)
    ws_tc.cell(row_num, 6, status)
    ws_tc.cell(row_num, 7, owner)
    ws_tc.cell(row_num, 8, approver)
    ws_tc.cell(row_num, 9, comments)
    ws_tc.cell(row_num, 10, ucs)
    ws_tc.cell(row_num, 11, frs)
    ws_tc.cell(row_num, 12, class_a)
    ws_tc.cell(row_num, 13, class_b)

    lines = tc_desc.count('\n') + max(1, len(tc_desc) // 45)
    ws_tc.row_dimensions[row_num].height = max(28, lines * 18)

# Row 27 là dòng hướng dẫn phân cách
ws_tc.cell(27, 1, "Insert lines above here only")
ws_tc.cell(27, 4, "Insert lines above here only")
ws_tc.cell(27, 6, "Insert lines above here only")
ws_tc.cell(27, 11, "Insert lines above here only")
ws_tc.row_dimensions[27].height = 20

# Dọn các dòng thừa từ 28 nếu có
for col in range(1, 14):
    ws_tc.cell(28, col, None)

tc_col_widths = {
    "A": 14, "B": 26, "C": 14, "D": 38, "E": 54, "F": 16,
    "G": 16, "H": 16, "I": 20, "J": 16, "K": 26, "L": 18, "M": 18
}
for col_let, w in tc_col_widths.items():
    ws_tc.column_dimensions[col_let].width = w


# =============================================================================
# 3. CẬP NHẬT SHEET 'Test Case Coverage Plan'
# =============================================================================
print("3. Đang cập nhật Sheet 'Test Case Coverage Plan'...")
ws_cov = wb['Test Case Coverage Plan']

for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
    cov_row = 3 + idx
    list_row = 2 + idx
    (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

    ws_cov.cell(cov_row, 1, f"='Test Case List'!C{list_row}")
    ws_cov.cell(cov_row, 2, f"='Test Case List'!D{list_row}")
    
    ws_cov.cell(cov_row, 3, iter_plan[0])
    ws_cov.cell(cov_row, 4, iter_plan[1])
    ws_cov.cell(cov_row, 5, iter_plan[2])
    ws_cov.cell(cov_row, 6, iter_plan[3])
    ws_cov.cell(cov_row, 7, iter_plan[4])

    ws_cov.cell(cov_row, 8, lead_tester)
    ws_cov.cell(cov_row, 9, "")
    ws_cov.row_dimensions[cov_row].height = 24

# Row 28 là dòng phân cách
ws_cov.cell(28, 1, "Insert Lines Above Here Only")
ws_cov.cell(28, 2, None)
ws_cov.cell(28, 3, "Insert Lines Above Here Only")
for col in range(4, 10):
    ws_cov.cell(28, col, None)
ws_cov.row_dimensions[28].height = 20

# Dọn dòng 29 nếu có
for col in range(1, 10):
    ws_cov.cell(29, col, None)

cov_col_widths = {"A": 14, "B": 38, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 18, "I": 20}
for col_let, w in cov_col_widths.items():
    ws_cov.column_dimensions[col_let].width = w


# =============================================================================
# 4. CẬP NHẬT 5 SHEETS 'Test Iteration 1' ĐẾN 'Test Iteration 5'
# =============================================================================
iteration_names = [
    "Iteration One",
    "Iteration Two",
    "Iteration Three",
    "Iteration Four",
    "Iteration Five"
]

for iter_idx in range(1, 6):
    sname = f"Test Iteration {iter_idx}"
    print(f"4.{iter_idx}. Đang cập nhật Sheet '{sname}'...")
    ws_it = wb[sname]

    # Đặt tên Iteration tại ô G31
    ws_it['G31'] = iteration_names[iter_idx - 1]

    for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
        r = 2 + idx
        list_row = 2 + idx
        iter_col_idx = 2 + iter_idx
        (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

        in_this_iter = iter_plan[iter_idx - 1]

        ws_it.cell(r, 1, f"='Test Case List'!C{list_row}")
        ws_it.cell(r, 2, f"='Test Case List'!D{list_row}")
        ws_it.cell(r, 3, f"=VLOOKUP(A{r},'Test Case Coverage Plan'!$A$3:$G$27,{iter_col_idx},FALSE)")
        ws_it.cell(r, 4, lead_tester if in_this_iter == "Yes" else None)

        # Trạng thái: 'Not Run' nếu trong đợt test, 'N/A' nếu không trong đợt test
        ws_it.cell(r, 5, "Not Run" if in_this_iter == "Yes" else "N/A")

        ws_it.cell(r, 6, "")
        ws_it.cell(r, 7, "")
        ws_it.cell(r, 8, "")
        ws_it.row_dimensions[r].height = 24

    # Row 27 là dòng phân cách
    ws_it.cell(27, 1, "Insert Rows Only Above Here")
    ws_it.cell(27, 2, None)
    ws_it.cell(27, 3, None)
    ws_it.cell(27, 4, "Insert Rows Only Above Here")
    for col in range(5, 9):
        ws_it.cell(27, col, None)
    ws_it.row_dimensions[27].height = 20

    # Dọn dòng 28
    for col in range(1, 9):
        ws_it.cell(28, col, None)

    it_col_widths = {"A": 14, "B": 38, "C": 22, "D": 16, "E": 18, "F": 20, "G": 20, "H": 24}
    for col_let, w in it_col_widths.items():
        ws_it.column_dimensions[col_let].width = w


# =============================================================================
# 5. CẬP NHẬT SHEET 'Stats'
# =============================================================================
print("5. Đang cập nhật Sheet 'Stats' (Chính xác 125 dòng data tương ứng 25 TC x 5 Iterations)...")
ws_stats = wb['Stats']

for iter_idx in range(1, 6):
    iter_sheet_name = f"Test Iteration {iter_idx}"
    start_stat_row = 2 + (iter_idx - 1) * 25

    for idx in range(25):
        stat_r = start_stat_row + idx
        source_r = 2 + idx
        list_r = 2 + idx

        ws_stats.cell(stat_r, 1, f"='{iter_sheet_name}'!A{source_r}")
        ws_stats.cell(stat_r, 2, f"='{iter_sheet_name}'!B{source_r}")
        ws_stats.cell(stat_r, 3, f"='{iter_sheet_name}'!$G$31")
        ws_stats.cell(stat_r, 4, f"='{iter_sheet_name}'!C{source_r}")
        ws_stats.cell(stat_r, 5, f"='{iter_sheet_name}'!D{source_r}")
        ws_stats.cell(stat_r, 6, f"='{iter_sheet_name}'!E{source_r}")
        ws_stats.cell(stat_r, 7, f"='Test Case List'!L{list_r}")
        ws_stats.cell(stat_r, 8, f"='Test Case List'!M{list_r}")
        ws_stats.row_dimensions[stat_r].height = 20

# Row 127 là dòng phân cách
ws_stats.cell(127, 1, "Insert Rows Only Above Here")
ws_stats.cell(127, 5, "Insert Rows Only Above Here")
ws_stats.row_dimensions[127].height = 20

# Dọn từ dòng 128 trở đi nếu có
for r_clear in range(128, ws_stats.max_row + 1):
    for c_clear in range(1, 9):
        ws_stats.cell(r_clear, c_clear, None)

stat_col_widths = {"A": 14, "B": 38, "C": 18, "D": 20, "E": 16, "F": 16, "G": 18, "H": 18}
for col_let, w in stat_col_widths.items():
    ws_stats.column_dimensions[col_let].width = w


# =============================================================================
# 6. ĐỒNG BỘ 100% FONT CHỮ TIMES NEW ROMAN & BẬT LƯỚI GRIDLINES
# =============================================================================
print("6. Đang áp dụng 100% font Times New Roman & Căn chỉnh wrap text cho toàn bộ file...")

for sname in wb.sheetnames:
    ws = wb[sname]
    if ws.views and ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                apply_tnr_font(cell, base_font=cell.font)

                if isinstance(cell.value, str) and ('\n' in cell.value or len(cell.value) > 25):
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal if cell.alignment else "left",
                        vertical="top" if '\n' in cell.value else (cell.alignment.vertical if cell.alignment else "center"),
                        wrap_text=True
                    )


# =============================================================================
# 7. LƯU WORKBOOK
# =============================================================================
print(f"7. Đang lưu file chuẩn hóa vào: {TEMPLATE_PATH}...")
wb.save(TEMPLATE_PATH)
print(">>> CẬP NHẬT TEST PLAN THEO CHUẨN CỦA CÔ GIÁO THÀNH CÔNG! <<<")
