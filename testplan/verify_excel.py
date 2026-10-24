# -*- coding: utf-8 -*-
import sys
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl

wb = openpyxl.load_workbook(r"testplan/Ke_Hoach_Va_Kich_Ban_Kiem_Thu_OfficePet.xlsx")
print("Sheets:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"Sheet {sname}: max_row={ws.max_row}, max_column={ws.max_column}")
    non_tnr = 0
    total_cells = 0
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None:
                total_cells += 1
                if cell.font and cell.font.name != "Times New Roman":
                    non_tnr += 1
    print(f"  -> Total valued cells: {total_cells}, Non-Times-New-Roman: {non_tnr}")

ws2 = wb["2. Danh_Sach_Test_Cases"]
statuses = set()
actual_non_empty = 0
dates_non_empty = 0
testers_non_empty = 0
defects_non_empty = 0

for r in range(6, ws2.max_row + 1):
    status = ws2.cell(row=r, column=10).value
    statuses.add(status)
    if ws2.cell(row=r, column=11).value:
        actual_non_empty += 1
    if ws2.cell(row=r, column=12).value:
        dates_non_empty += 1
    if ws2.cell(row=r, column=13).value:
        testers_non_empty += 1
    if ws2.cell(row=r, column=14).value:
        defects_non_empty += 1

print("Statuses found in Col 10:", statuses)
print(f"Actual results filled (Col 11): {actual_non_empty}")
print(f"Dates filled (Col 12): {dates_non_empty}")
print(f"Testers filled (Col 13): {testers_non_empty}")
print(f"Defects filled (Col 14): {defects_non_empty}")
