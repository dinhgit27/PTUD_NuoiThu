# -*- coding: utf-8 -*-
"""
Script cập nhật chính xác theo đúng hình ảnh người dùng cung cấp:
- Thành viên 1 (Nhóm trưởng): Nguyễn Đỉnh - MSSV: 123000431
  -> Phụ trách: Auth, Desktop PiP, Admin Console, Security & NFR (TS01, TS03, TS09, TS10)
  -> Approver: Nguyễn Đỉnh (123000431) - Lead QA
- Thành viên 2: Hồ Tâm Phi Hùng - MSSV: 123000287
  -> Phụ trách: Nuôi thú 11 hoạt ảnh, Nhiệm vụ, Shop & Kho đồ (TS02, TS06, TS08)
- Thành viên 3: Võ Lê Hoàng Nhã - MSSV: 123000311
  -> Phụ trách: Pomodoro, Wellness Lofi/Thở, Chatbot AI, Vườn chung (TS04, TS05, TS07)

Áp dụng đồng bộ trên tất cả các trang:
1. ReadMe
2. Test Case List (Owner, Approver, Lookup table)
3. Test Case Coverage Plan (Planned Lead Tester, Lookup table)
4. Test Iteration 1 đến 5 (Tester, Lookup table)
5. Stats (Liên kết tự động)
6. Pivot Tables
"""

import sys
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FILE_PATH = r"c:\Users\PHI HUNG\Downloads\PTUD_NuoiThu\testplan\test-plan-iterations.xlsx"

# Đúng chuẩn theo hình ảnh người dùng cung cấp:
LEADER_DINH = "Nguyễn Đỉnh (123000431)"
MEMBER_HUNG = "Hồ Tâm Phi Hùng (123000287)"
MEMBER_NHA  = "Võ Lê Hoàng Nhã (123000311)"
APPROVER_NAME = "Nguyễn Đỉnh (123000431) - Lead QA"

# 25 Kịch bản kiểm thử phân bổ đúng người đúng việc theo hình:
# - Nguyễn Đỉnh: TS01 (Auth), TS03 (Desktop PiP), TS09 (Admin Console), TS10 (NFR & Security)
# - Hồ Tâm Phi Hùng: TS02 (Nuôi thú 11 hoạt ảnh), TS06 (Nhiệm vụ), TS08 (Shop & Kho đồ)
# - Võ Lê Hoàng Nhã: TS04 (Pomodoro & Wellness), TS05 (Chatbot AI), TS07 (Shared Garden & Bạn bè)

OFFICEPET_25_TEST_CASES = [
    # --- PHÂN HỆ 01: XÁC THỰC & QUẢN LÝ TÀI KHOẢN (TS01) - Nguyễn Đỉnh phụ trách ---
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.1",
        "Đăng ký tài khoản mới hợp lệ & Kiểm thực lỗi email/mật khẩu",
        "[Mục đích] Xác nhận quy trình đăng ký tài khoản nhân viên mới vào hệ thống.\n[Điều kiện] Người dùng chưa có tài khoản, đang ở màn hình Đăng ký.\n[Các bước] 1. Nhập Họ tên 'Nguyễn Văn A', Email hợp lệ, Mật khẩu >= 8 ký tự, chọn Phòng ban.\n2. Nhấn nút 'Đăng ký'.\n3. Thử lại với email trùng hoặc mật khẩu < 8 ký tự.\n[Kỳ vọng] Hệ thống tạo tài khoản mã hóa bcrypt, báo thành công; chặn email trùng và mật khẩu yếu.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-01",
        "FR-01.1, FR-01.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.2",
        "Đăng nhập phân quyền vai trò (Role ADMIN & USER) và JWT Token",
        "[Mục đích] Xác thực cơ chế phân quyền truy cập giữa Quản trị viên và Nhân viên.\n[Điều kiện] Đã có sẵn tài khoản Admin và User trong hệ thống.\n[Các bước] 1. Đăng nhập User -> điều hướng vào Vườn thú cưng làm việc.\n2. Đăng xuất tài khoản.\n3. Đăng nhập Admin -> hiển thị menu Bảng điều khiển Quản trị (Admin Console).\n[Kỳ vọng] JWT token cấp đúng role; User bị chặn vào /admin, Admin truy cập được trang quản trị.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-01",
        "FR-01.3, FR-01.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS01",
        "Xác thực & Tài khoản",
        "TC1.3",
        "Quên mật khẩu qua mã OTP Email & Cập nhật ảnh đại diện Avatar",
        "[Mục đích] Kiểm tra quy trình khôi phục mật khẩu qua OTP và cập nhật thông tin cá nhân.\n[Điều kiện] Tài khoản người dùng có email hợp lệ.\n[Các bước] 1. Bấm 'Quên mật khẩu', nhập email nhận mã OTP 6 chữ số.\n2. Nhập mã OTP và thiết lập mật khẩu mới.\n3. Đăng nhập đổi ảnh đại diện trong trang Hồ sơ cá nhân.\n[Kỳ vọng] OTP gửi thành công; cập nhật mật khẩu và avatar hiển thị tức thì trên giao diện.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-01",
        "FR-01.5, FR-01.6",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),

    # --- PHÂN HỆ 02: NUÔI & TƯƠNG TÁC THÚ CƯNG 11 HOẠT ẢNH (TS02) - Hồ Tâm Phi Hùng phụ trách ---
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.1",
        "Khởi tạo nhận nuôi thú cưng & Cơ chế suy giảm chỉ số sinh học tự nhiên",
        "[Mục đích] Xác thực nhận nuôi thú cưng ban đầu và suy giảm chỉ số Đói/Khát/Năng lượng theo thời gian.\n[Điều kiện] Tài khoản mới đăng nhập lần đầu vào hệ thống.\n[Các bước] 1. Chọn thú cưng mặc định (Corgi), đặt tên 'Bắp Cải', nhận nuôi.\n2. Quan sát chỉ số Đói và Khát giảm dần sau mỗi chu kỳ 5-10 phút.\n[Kỳ vọng] Bản ghi UserPet tạo đúng Level 1, 100 Năng lượng, 80 Đói/Khát; thanh trạng thái đổi màu vàng/đỏ khi < 30%.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-02",
        "FR-02.1, FR-02.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.2",
        "Hành động Cho ăn (EATING) và Cho uống nước (DRINKING) tách bạch",
        "[Mục đích] Kiểm tra hoạt ảnh chuẩn hóa riêng biệt giữa Ăn và Uống theo chuẩn 11 hoạt ảnh mới.\n[Điều kiện] Thú cưng có chỉ số Đói < 80 và Khát < 80.\n[Các bước] 1. Nhấn nút 'Cho ăn' -> quan sát hoạt ảnh EATING.\n2. Nhấn nút 'Cho uống nước' -> quan sát hoạt ảnh DRINKING tách biệt riêng.\n[Kỳ vọng] Cho ăn hồi +30 Đói với hoạt ảnh nhai thức ăn; Cho uống hồi +25 Khát với hoạt ảnh uống nước riêng; kết thúc trở về IDLE.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-02",
        "FR-02.3, FR-02.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.3",
        "Tương tác vuốt ve (PETTING) giới hạn +30 EXP/ngày và Hoạt ảnh HAPPY / SAD",
        "[Mục đích] Kiểm tra tương tác vuốt ve, cơ chế chống spam cày cấp và chuyển đổi hoạt ảnh cảm xúc.\n[Điều kiện] Thú cưng đang đứng yên ở trạng thái IDLE.\n[Các bước] 1. Click chuột vuốt ve thú cưng liên tục 6 lần (+5 EXP/lần).\n2. Click tiếp tục lần 7 trở đi trong ngày.\n3. Để chỉ số đói khát tụt xuống < 15%.\n[Kỳ vọng] Thú cưng phát tim (PETTING), cộng tối đa +30 EXP/ngày; click tiếp vẫn vui nhưng không cộng thêm EXP; khi đói hiển thị HUNGRY/SAD.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-02",
        "FR-02.5, FR-02.6",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS02",
        "Nuôi & Chăm sóc Thú Cưng",
        "TC2.4",
        "Cho thú cưng ngủ (SLEEPING) hồi năng lượng và Thăng cấp thú cưng",
        "[Mục đích] Kiểm tra cơ chế hồi phục năng lượng khi ngủ và thăng cấp (Level Up).\n[Điều kiện] Năng lượng < 50, điểm kinh nghiệm đạt 95/100 EXP.\n[Các bước] 1. Bấm 'Đi ngủ' -> phòng tối, hoạt ảnh SLEEPING phập phồng chữ Zzz.\n2. Bấm 'Đánh thức'.\n3. Thực hiện hành động chăm sóc để đạt mốc 100 EXP.\n[Kỳ vọng] Năng lượng tăng dần khi ngủ; khi đạt 100 EXP kích hoạt Level Up, hoạt ảnh nhảy múa DANCING và nhận thưởng mở khóa.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-02",
        "FR-02.7, FR-02.8",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),

    # --- PHÂN HỆ 03: DESKTOP COMPANION & PIP (TS03) - Nguyễn Đỉnh phụ trách ---
    (
        "TS03",
        "Desktop Companion & PiP",
        "TC3.1",
        "Thu nhỏ thú cưng ra màn hình Desktop (PiP) và Kéo thả ghi nhớ tọa độ",
        "[Mục đích] Kiểm tra tính năng Document PiP luôn nổi trên desktop và ghi nhớ tọa độ màn hình.\n[Điều kiện] Trình duyệt có hỗ trợ Document PiP API.\n[Các bước] 1. Nhấn nút Desktop PiP trên web.\n2. Kéo thả thú cưng sang góc phải dưới màn hình.\n3. Đóng cửa sổ và mở lại PiP.\n[Kỳ vọng] Cửa sổ popup trong suốt nổi trên màn hình; khi mở lại tự động khôi phục đúng tọa độ đã kéo thả trước đó.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-03",
        "FR-03.1, FR-03.2",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS03",
        "Desktop Companion & PiP",
        "TC3.2",
        "Thú cưng tự do đi lại trong ranh giới (Roaming) và Menu thao tác nhanh",
        "[Mục đích] Kiểm tra chuyển động đi bộ tự nhiên và bảng menu chuột phải trên desktop.\n[Điều kiện] Cửa sổ Desktop PiP đang mở.\n[Các bước] 1. Bật chế độ Roaming, quan sát thú đi dạo.\n2. Nhấp chuột phải (Right-click) lên thú cưng.\n3. Chọn nút 'Cho ăn nhanh'.\n[Kỳ vọng] Thú cưng tự động quay đầu khi chạm mép màn hình (WALKING); menu chuột phải hiển thị mượt mà và thực hiện hành động tức thì.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-03",
        "FR-03.3, FR-03.4",
        "Minor",
        "P3 - Low",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),

    # --- PHÂN HỆ 04: ĐỒNG HÀNH NĂNG SUẤT POMODORO & WELLNESS (TS04) - Võ Lê Hoàng Nhã phụ trách ---
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.1",
        "Phiên Pomodoro 25 phút, thú cưng gõ phím TYPING và Nhận thưởng hoàn thành",
        "[Mục đích] Kiểm tra đồng hồ đếm giờ Pomodoro, đồng bộ trạng thái gõ phím và cộng thưởng hoàn thành.\n[Điều kiện] Thú cưng đang ở trạng thái IDLE.\n[Các bước] 1. Bắt đầu phiên 25 phút -> quan sát thú chuyển sang TYPING.\n2. Tạm dừng -> trở về IDLE.\n3. Để đếm hết 25 phút.\n[Kỳ vọng] Đồng hồ chạy chuẩn xác; thú cưng ngồi gõ máy tính chăm chỉ; kết thúc phiên nhận +50 EXP, +20 Vàng và phát chuông báo.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-04",
        "FR-04.1, FR-04.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.2",
        "Chế độ Nghiêm khắc cảnh báo xao nhãng khi người dùng rời tab làm việc",
        "[Mục đích] Kiểm tra tính năng chống xao nhãng của Pomodoro trong giờ làm việc.\n[Điều kiện] Bật 'Chế độ Nghiêm khắc' (Strict Mode) trong cài đặt Pomodoro.\n[Các bước] 1. Bắt đầu phiên làm việc.\n2. Chuyển sang tab khác (Youtube/MXH) hơn 15 giây.\n3. Quay trở lại ứng dụng.\n[Kỳ vọng] Trình duyệt gửi thông báo đẩy nhắc nhở tập trung; thú cưng chuyển vẻ mặt buồn (SAD) cảnh báo người dùng.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-04",
        "FR-04.3",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),
    (
        "TS04",
        "Pomodoro & Năng suất",
        "TC4.3",
        "Thực hiện Bài tập thở thư giãn (4-7-8 / Box) và Trình phát nhạc Lofi",
        "[Mục đích] Kiểm tra các tính năng chăm sóc sức khỏe tinh thần và phát nhạc nền tập trung.\n[Điều kiện] Thiết bị có loa / tai nghe kết nối hoạt động.\n[Các bước] 1. Chọn kỹ thuật thở 4-7-8, làm theo nhịp thở trực quan trên màn hình.\n2. Bật trình phát nhạc Lofi, chỉnh âm lượng, chuyển sang track tiếng mưa.\n[Kỳ vọng] Đồ họa vòng tròn thở co giãn đúng nhịp giây; nhạc Lofi phát êm ái, chuyển bài mượt mà không lỗi âm thanh.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-04",
        "FR-04.4, FR-04.5",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),

    # --- PHÂN HỆ 05: CHATBOT AI TRÒ CHUYỆN ĐỒNG HÀNH (TS05) - Võ Lê Hoàng Nhã phụ trách ---
    (
        "TS05",
        "Chatbot AI Đồng hành",
        "TC5.1",
        "Hộp thoại Chatbot AI, gửi câu hỏi công việc và Nhận phản hồi thông minh",
        "[Mục đích] Kiểm tra giao diện ngăn kéo chat và phản hồi thông minh từ trợ lý AI.\n[Điều kiện] Kết nối Internet hoạt động, API Key được cấu hình.\n[Các bước] 1. Mở ngăn kéo chat bên cạnh thú cưng.\n2. Nhập câu hỏi xin lời khuyên sắp xếp công việc.\n3. Nhấn gửi tin nhắn.\n[Kỳ vọng] Tin nhắn hiển thị đẹp mắt, thú cưng hiển thị bong bóng suy nghĩ và trả lời động viên khoa học trong 1-3 giây.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-05",
        "FR-05.1, FR-05.2",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),
    (
        "TS05",
        "Chatbot AI Đồng hành",
        "TC5.2",
        "Đồng bộ hoạt ảnh biểu cảm thú cưng (HAPPY / SAD) theo sắc thái trò chuyện",
        "[Mục đích] Kiểm tra sự gắn kết cảm xúc giữa nội dung trò chuyện AI và hoạt ảnh thú cưng.\n[Điều kiện] Chatbot AI đang hoạt động bình thường.\n[Các bước] 1. Gửi tin nhắn khoe tin vui hoàn thành dự án -> quan sát thú cưng.\n2. Gửi tin nhắn tâm sự chuyện buồn áp lực -> quan sát thú cưng.\n[Kỳ vọng] Với tin vui thú cưng kích hoạt HAPPY / DANCING; với chuyện buồn thú làm động tác vẫy tay chia sẻ an ủi.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-05",
        "FR-05.3",
        "Minor",
        "P3 - Low",
        ("No", "No", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),

    # --- PHÂN HỆ 06: NHIỆM VỤ & THI ĐUA PHÒNG BAN (TS06) - Hồ Tâm Phi Hùng phụ trách ---
    (
        "TS06",
        "Nhiệm vụ & Thi đua",
        "TC6.1",
        "Nhiệm vụ hàng ngày, tự động cộng dồn tiến độ và Nhận thưởng EXP",
        "[Mục đích] Kiểm tra hệ thống tự động ghi nhận hành vi người dùng vào tiến độ nhiệm vụ.\n[Điều kiện] Nhiệm vụ 'Cho ăn 3 lần' đang có tiến độ 0/3.\n[Các bước] 1. Thực hiện cho thú ăn đủ 3 lần.\n2. Mở tab Nhiệm vụ kiểm tra tiến độ.\n3. Bấm nút 'Nhận thưởng'.\n[Kỳ vọng] Thanh tiến trình tự động nhảy 3/3; nút 'Nhận thưởng' sáng lên; bấm nhận cộng đúng EXP và nút chuyển sang 'Đã nhận'.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-06",
        "FR-06.1, FR-06.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS06",
        "Nhiệm vụ & Thi đua",
        "TC6.2",
        "Bảng xếp hạng thi đua cá nhân, phòng ban và Cơ chế reset tự động lúc 00:00",
        "[Mục đích] Kiểm tra bảng xếp hạng thi đua và cơ chế reset ngày mới của hệ thống.\n[Điều kiện] Đã có dữ liệu tích lũy điểm của các phòng ban.\n[Các bước] 1. Mở Bảng xếp hạng, chuyển đổi bộ lọc Tuần/Tháng/Phòng ban.\n2. Kiểm tra reset nhiệm vụ sau mốc 00:00.\n[Kỳ vọng] Danh sách xếp hạng hiển thị đúng theo điểm số giảm dần; nhiệm vụ ngày cũ được làm mới tự động lúc 00:00.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-06",
        "FR-06.3, FR-06.4",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),

    # --- PHÂN HỆ 07: KHU VƯỜN CỘNG ĐỒNG & HỆ THỐNG BẠN BÈ (TS07) - Võ Lê Hoàng Nhã phụ trách ---
    (
        "TS07",
        "Khu vườn & Bạn bè",
        "TC7.1",
        "Khu vườn chung đa thú cưng phòng ban, Kết nối bạn bè & Ghé thăm vườn",
        "[Mục đích] Kiểm tra render nhiều thú cưng trên Canvas chung, tính năng kết bạn và thăm vườn.\n[Điều kiện] Đồng nghiệp trong cùng phòng ban cùng trực tuyến.\n[Các bước] 1. Mở 'Vườn Chung' xem thú cưng các đồng nghiệp dạo chơi.\n2. Tìm kiếm email gửi kết bạn.\n3. Ghé thăm vườn riêng của người bạn đó.\n[Kỳ vọng] Thú cưng hiển thị đồng thời mượt mà không đè tọa độ; kết bạn thành công và chuyển cảnh thăm vườn chuẩn xác.",
        "Approved",
        MEMBER_NHA,
        APPROVER_NAME,
        "",
        "UC-07",
        "FR-07.1, FR-07.2",
        "Major",
        "P2 - Medium",
        ("No", "No", "Yes", "Yes", "Yes"),
        MEMBER_NHA
    ),

    # --- PHÂN HỆ 08: CỬA HÀNG SHOP & KHO ĐỒ NỘI BỘ (TS08) - Hồ Tâm Phi Hùng phụ trách ---
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.1",
        "Mua Thú cưng mới từ Shop bằng điểm EXP và Đổi thú cưng đại diện trong Kho đồ",
        "[Mục đích] Kiểm tra giao dịch mua thú cưng bằng điểm tích lũy và chuyển đổi thú cưng hoạt động.\n[Điều kiện] Tài khoản có đủ 500 EXP.\n[Các bước] 1. Vào Shop tab Thú cưng, chọn 'Cáo Tuyết' (500 EXP), nhấn Mua.\n2. Vào Kho đồ, bấm 'Đặt làm thú cưng đại diện'.\n[Kỳ vọng] Số dư EXP bị trừ đúng 500; Cáo Tuyết được thêm vào UserPet; Màn hình chính và Desktop đổi ngay sang Cáo Tuyết.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-08",
        "FR-08.1, FR-08.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.2",
        "Mua Skin/Trang phục và Trang bị (Equip/Unequip) lên thú cưng",
        "[Mục đích] Kiểm tra tính năng mua sắm trang phục và hiển thị skin layer trên canvas.\n[Điều kiện] Người dùng có đủ Vàng và sở hữu thú cưng tương thích.\n[Các bước] 1. Mua skin 'Áo Hoodie Dev'.\n2. Vào Kho đồ nhấn 'Trang bị' (Equip).\n3. Bấm 'Tháo ra' (Unequip).\n[Kỳ vọng] Khi Equip thú cưng hiển thị áo hoodie mượt mà; khi Unequip trở về diện mạo nguyên bản mặc định.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-08",
        "FR-08.3, FR-08.4",
        "Major",
        "P2 - Medium",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),
    (
        "TS08",
        "Cửa hàng Shop & Kho đồ",
        "TC8.3",
        "Mua bối cảnh Nền vườn mới và Áp dụng thay đổi không gian làm việc",
        "[Mục đích] Kiểm tra tùy biến bối cảnh nền và hiệu ứng gia tăng EXP nội tại từ vật phẩm.\n[Điều kiện] Đủ điểm EXP mua background mới trong Shop.\n[Các bước] 1. Mua nền 'Văn phòng Cyberpunk'.\n2. Vào Kho đồ bấm 'Áp dụng làm nền vườn'.\n[Kỳ vọng] Khu vườn đổi ngay sang phong cách Cyberpunk, thú cưng đứng đúng vị trí sàn; buff EXP hoạt động chính xác.",
        "Approved",
        MEMBER_HUNG,
        APPROVER_NAME,
        "",
        "UC-08",
        "FR-08.5",
        "Minor",
        "P3 - Low",
        ("No", "Yes", "Yes", "Yes", "Yes"),
        MEMBER_HUNG
    ),

    # --- PHÂN HỆ 09: CỔNG QUẢN TRỊ ADMIN CONSOLE & FORGE (TS09) - Nguyễn Đỉnh phụ trách ---
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.1",
        "Bảng điều khiển số liệu hệ thống và Đổi quyền người dùng (Role USER <-> ADMIN)",
        "[Mục đích] Kiểm tra các widget thống kê quản trị và phân quyền bảo mật tài khoản.\n[Điều kiện] Đăng nhập bằng tài khoản Quản trị viên (Admin).\n[Các bước] 1. Mở Admin Dashboard -> xem số liệu tổng người dùng, thú cưng.\n2. Vào Quản lý Người dùng, đổi role một tài khoản thử nghiệm.\n[Kỳ vọng] Số liệu tải chính xác từ DB; role được cập nhật ngay lập tức; chặn tài khoản bị ban đăng nhập.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-09",
        "FR-09.1, FR-09.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.2",
        "Xưởng Pet Forge tạo thú 11 hành động, bỏ ô mô tả thừa, dọn sạch form và Hiện ở Shop",
        "[Mục đích] Xác thực cải tiến Pet Forge: preview đủ 11 hành động, không có ô mô tả thừa, dọn sạch UI sau tạo và lưu DB hiện ngay ở Shop.\n[Điều kiện] Admin đang ở công cụ Pet Forge.\n[Các bước] 1. Upload Sprite sheet 11 hàng.\n2. Xác nhận form không có ô 'Mô tả'.\n3. Nhập Tên, Giá và nhấn 'Tạo thú cưng'.\n4. Kiểm tra form reset sạch sẽ và sang Quản lý Shop kiểm tra.\n[Kỳ vọng] Form dọn sạch giao diện ngay sau khi tạo; thú mới lưu vào bảng PetType và xuất hiện ngay lập tức trong Quản lý Shop.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-09",
        "FR-09.3, FR-09.4",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS09",
        "Cổng Quản trị Admin",
        "TC9.3",
        "Xưởng tạo Skin, Xưởng tạo Vật phẩm hoạt ảnh và Chức năng Xóa thú cưng khỏi hệ thống",
        "[Mục đích] Xác thực việc tách riêng 2 xưởng tạo Skin/Vật phẩm, loại bỏ mục vật phẩm tĩnh cũ và kiểm tra nút xóa thú.\n[Điều kiện] Admin ở trang Quản lý Shop.\n[Các bước] 1. Kiểm tra menu điều hướng sang Tạo Skin và Tạo Vật phẩm (đã bỏ nút thêm vật phẩm cũ và mục vật phẩm tĩnh).\n2. Nhấn nút Xóa một thú cưng thử nghiệm.\n[Kỳ vọng] Modal xác nhận xóa hiển thị an toàn; sau khi bấm xóa, thú cưng bị gỡ khỏi danh sách quản lý và database.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-09",
        "FR-09.5, FR-09.6",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),

    # --- PHÂN HỆ 10: PHI CHỨC NĂNG, BẢO MẬT & TƯƠNG THÍCH NGƯỢC (TS10) - Nguyễn Đỉnh phụ trách ---
    (
        "TS10",
        "Phi Chức Năng & Bảo Mật",
        "TC10.1",
        "Tương thích ngược thú cũ 17 hàng và Fallback CELEBRATING->HAPPY, EXERCISING->DANCING",
        "[Mục đích] Đảm bảo các con thú tạo trước đây (chuẩn 17 hàng) và thú mới (11 hàng) đều hiển thị bình thường không bị lỗi.\n[Điều kiện] Hệ thống có thú cưng sprite sheet 17 hàng cũ.\n[Các bước] 1. Mở xem thú cưng cũ 17 hàng.\n2. Kích hoạt các sự kiện cũ CELEBRATING và EXERCISING.\n[Kỳ vọng] Thú cũ chạy mượt mà không lệch tọa độ frame; cơ chế fallback tự động chuyển CELEBRATING->HAPPY và EXERCISING->DANCING không gây crash console.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-10",
        "FR-10.1, FR-10.2",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
    (
        "TS10",
        "Phi Chức Năng & Bảo Mật",
        "TC10.2",
        "Render Canvas 60 FPS không rò rỉ RAM, Chặn Route Guard /admin và Chống XSS",
        "[Mục đích] Kiểm tra hiệu năng mượt mà của Canvas và kiểm thử an toàn bảo mật hệ thống.\n[Điều kiện] Đăng nhập bằng tài khoản User thông thường.\n[Các bước] 1. Đo FPS khi thú cưng chuyển động trong 30 phút.\n2. Nhập URL /admin/pets trên trình duyệt.\n3. Thử inject payload <script>alert(1)</script> vào tên thú.\n[Kỳ vọng] FPS đạt 55-60 FPS ổn định, RAM không phình; route /admin bị chặn 403; payload XSS bị escape an toàn.",
        "Approved",
        LEADER_DINH,
        APPROVER_NAME,
        "",
        "UC-10",
        "FR-10.3, FR-10.4, FR-10.5",
        "Critical",
        "P1 - High",
        ("Yes", "Yes", "Yes", "Yes", "Yes"),
        LEADER_DINH
    ),
]

assert len(OFFICEPET_25_TEST_CASES) == 25, "Bắt buộc chính xác 25 test cases"

print(f"Đang mở file: {FILE_PATH}...")
wb = openpyxl.load_workbook(FILE_PATH)

FONT_TNR = "Times New Roman"

def set_tnr(cell, size=None, bold=None, italic=None, color=None):
    base_f = cell.font
    cell.font = Font(
        name=FONT_TNR,
        size=size if size is not None else (base_f.size if base_f else 11),
        bold=bold if bold is not None else (base_f.bold if base_f else False),
        italic=italic if italic is not None else (base_f.italic if base_f else False),
        color=color if color is not None else (base_f.color if base_f else None)
    )

# =============================================================================
# 1. SHEET 'ReadMe': CẬP NHẬT ĐÚNG HÌNH ẢNH NGƯỜI DÙNG CUNG CẤP
# =============================================================================
print("1. Đang cập nhật Sheet 'ReadMe' theo đúng hình ảnh...")
ws_readme = wb['ReadMe']

# Xóa các giá trị cũ
for r in range(1, 30):
    for c in range(1, 10):
        ws_readme.cell(r, c, None)

readme_content = [
    (1, "A", "Tên Đồ Án:", True, 12),
    (1, "B", "HỆ THỐNG NUÔI THÚ CƯNG ẢO ĐỒNG HÀNH NĂNG SUẤT VĂN PHÒNG (OFFICEPET - PTUD_NUOITHU)", True, 13),
    (3, "A", "Nhóm Thực Hiện:", True, 11),
    (3, "B", "NHÓM SINH VIÊN THỰC HIỆN ĐỀ TÀI PHÁT TRIỂN ỨNG DỤNG", True, 11),
    # ĐÚNG NHƯ HÌNH:
    (4, "A", "Thành viên 1 (Nhóm trưởng):", True, 11),
    (4, "B", "Nguyễn Đỉnh - MSSV: 123000431 (Phụ trách: Auth, Desktop PiP, Admin Console, Security & NFR)", False, 11),
    (5, "A", "Thành viên 2:", True, 11),
    (5, "B", "Hồ Tâm Phi Hùng  - MSSV: 123000287 (Phụ trách: Nuôi thú 11 hoạt ảnh, Nhiệm vụ, Shop & Kho đồ)", False, 11),
    (6, "A", "Thành viên 3:", True, 11),
    (6, "B", "Võ Lê Hoàng Nhã - MSSV: 123000311 (Phụ trách: Pomodoro, Wellness Lofi/Thở, Chatbot AI, Vườn chung)", False, 11),
    (8, "A", "Phiên bản Hệ thống:", True, 11),
    (8, "B", "Release v1.0.0 (Chuẩn hóa 11 Hoạt ảnh Thú Cưng & Kiến trúc Tương thích ngược)", False, 11),
    (9, "A", "Ngày lập Kế hoạch:", True, 11),
    (9, "B", "24/10/2026", False, 11),
    (10, "A", "Bản quyền Đề tài:", True, 11),
    (10, "B", "OfficePet System - PTUD_NuoiThu © 2026", False, 11),
    (12, "A", "Quy Chuẩn Tài Liệu:", True, 11),
    (12, "B", "Kế hoạch kiểm thử phân kỳ (Test Plan Iterations Template) chuẩn mực theo yêu cầu của Giảng viên hướng dẫn.", False, 11),
    (13, "A", "Phạm Vi Nghiệp Vụ:", True, 11),
    (13, "B", "Bao phủ toàn diện 10 Phân hệ chức năng & phi chức năng (TS01 -> TS10) với 25 kịch bản tích hợp và 5 Vòng lặp kiểm thử (Iterations).", False, 11),
    (15, "A", "Kế Hoạch 5 Vòng Lặp (Iterations):", True, 12),
    (15, "B", "Phân kỳ kiểm thử lũy tiến:", True, 11),
    (16, "A", "• Vòng lặp 1 (Iteration One):", True, 11),
    (16, "B", "Smoke Test: Kiểm thử độ ổn định các luồng chức năng sống còn (Critical): Đăng ký, Đăng nhập, Nuôi thú 11 hoạt ảnh, Pomodoro, Admin Forge, Tương thích ngược 17 hàng.", False, 11),
    (17, "A", "• Vòng lặp 2 (Iteration Two):", True, 11),
    (17, "B", "Core Features: Kiểm thử sâu chức năng lõi: Quản lý hồ sơ, Desktop PiP kéo thả lưu vị trí, Chế độ nghiêm khắc Pomodoro, Bài tập thở, Nhạc Lofi, Kho đồ.", False, 11),
    (18, "A", "• Vòng lặp 3 (Iteration Three):", True, 11),
    (18, "B", "Full Functional: Kiểm thử toàn diện 100% chức năng: Cửa hàng Shop, Vườn chung phòng ban, Hệ thống kết nối bạn bè, Chatbot AI đồng hành cảm xúc.", False, 11),
    (19, "A", "• Vòng lặp 4 (Iteration Four):", True, 11),
    (19, "B", "Regression & NFR: Kiểm thử hồi quy, kiểm tra hiệu năng Canvas 60 FPS, Fallback hoạt ảnh cũ, Route Guard bảo vệ Admin và chống tấn công XSS.", False, 11),
    (20, "A", "• Vòng lặp 5 (Iteration Five):", True, 11),
    (20, "B", "Final Acceptance / UAT: Nghiệm thu toàn diện hệ thống trước khi đóng gói bàn giao.", False, 11),
    (22, "A", "Bảng Quy Ước Trạng Thái (Legend):", True, 12),
    (22, "B", "Quy định mã trạng thái kiểm thử trung thực (Không fake kết quả):", True, 11),
    (23, "A", "• Not Run:", True, 11),
    (23, "B", "Chưa kiểm thử - Sẵn sàng thực thi trong đợt test (Cam kết trung thực, không tạo kết quả ảo khi chưa bắt đầu test).", False, 11),
    (24, "A", "• Run - Passed:", True, 11),
    (24, "B", "Đã chạy test thực tế và kết quả thực tế trùng khớp hoàn toàn với kết quả mong đợi.", False, 11),
    (25, "A", "• Run - Failed:", True, 11),
    (25, "B", "Đã chạy test thực tế và phát sinh lỗi sai lệch so với yêu cầu kỹ thuật.", False, 11),
    (26, "A", "• Blocked:", True, 11),
    (26, "B", "Bị chặn không thể thực hiện kiểm thử do chức năng tiên quyết bị lỗi nghiêm trọng.", False, 11),
    (27, "A", "• N/A:", True, 11),
    (27, "B", "Test case không nằm trong phạm vi kế hoạch thực thi của đợt Iteration này.", False, 11),
]

for r, col, val, is_bold, f_size in readme_content:
    cell = ws_readme[f"{col}{r}"]
    cell.value = val
    cell.font = Font(name=FONT_TNR, size=f_size, bold=is_bold)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

for r in range(1, 30):
    if r in [1, 15, 22]:
        ws_readme.row_dimensions[r].height = 28
    elif r in [16, 17, 18, 19, 20]:
        ws_readme.row_dimensions[r].height = 36
    elif r in [23, 24, 25, 26, 27]:
        ws_readme.row_dimensions[r].height = 24
    else:
        ws_readme.row_dimensions[r].height = 22

ws_readme.column_dimensions["A"].width = 32
ws_readme.column_dimensions["B"].width = 96
ws_readme.column_dimensions["C"].width = 16


# =============================================================================
# 2. SHEET 'Test Case List'
# =============================================================================
print("2. Đang cập nhật Sheet 'Test Case List' (Đúng phân công Nguyễn Đỉnh, Hồ Tâm Phi Hùng, Võ Lê Hoàng Nhã)...")
ws_tc = wb['Test Case List']

# Lookup data validation
statuses_list = ['Not Started', 'In Progress', 'Draft', 'Pending Approval', 'Approved']
for i, s in enumerate(statuses_list):
    ws_tc.cell(row=31 + i, column=1, value=s)

class_a_list = ['Critical', 'Major', 'Minor', 'Trivial']
for i, ca in enumerate(class_a_list):
    ws_tc.cell(row=31 + i, column=4, value=ca)

class_b_list = ['P1 - High', 'P2 - Medium', 'P3 - Low', 'P4 - Lowest']
for i, cb in enumerate(class_b_list):
    ws_tc.cell(row=46 + i, column=4, value=cb)

# Cập nhật danh sách Owner:
owners_list = [LEADER_DINH, MEMBER_HUNG, MEMBER_NHA, APPROVER_NAME]
for i, ow in enumerate(owners_list):
    ws_tc.cell(row=42 + i, column=1, value=ow)

ws_tc.row_dimensions[1].height = 30
for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
    row_num = 2 + idx
    (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

    ws_tc.cell(row_num, 1, sc_id)
    ws_tc.cell(row_num, 2, sc_name)
    ws_tc.cell(row_num, 3, tc_id)
    ws_tc.cell(row_num, 4, tc_name)
    ws_tc.cell(row_num, 5, tc_desc)
    ws_tc.cell(row_num, 6, status)
    ws_tc.cell(row_num, 7, owner)
    ws_tc.cell(row_num, 8, approver)
    ws_tc.cell(row_num, 9, comments)
    ws_tc.cell(row_num, 10, ucs)
    ws_tc.cell(row_num, 11, frs)
    ws_tc.cell(row_num, 12, class_a)
    ws_tc.cell(row_num, 13, class_b)

    lines_desc = tc_desc.count('\n') + max(1, len(tc_desc) // 55)
    ws_tc.row_dimensions[row_num].height = max(75, lines_desc * 19)

ws_tc.cell(27, 1, "Insert lines above here only")
ws_tc.cell(27, 4, "Insert lines above here only")
ws_tc.cell(27, 6, "Insert lines above here only")
ws_tc.cell(27, 11, "Insert lines above here only")
ws_tc.row_dimensions[27].height = 22

for col in range(1, 15):
    ws_tc.cell(28, col, None)

tc_col_widths = {
    "A": 16,
    "B": 28,
    "C": 16,
    "D": 45,
    "E": 68,
    "F": 18,
    "G": 34, # Owner (Đủ rộng cho Tên + MSSV)
    "H": 36, # Approver (Đủ rộng cho Tên + MSSV - Lead QA)
    "I": 20,
    "J": 16,
    "K": 26,
    "L": 18,
    "M": 18
}
for col_let, w in tc_col_widths.items():
    ws_tc.column_dimensions[col_let].width = w


# =============================================================================
# 3. SHEET 'Test Case Coverage Plan'
# =============================================================================
print("3. Đang cập nhật Sheet 'Test Case Coverage Plan' (Planned Lead Tester theo phân công)...")
ws_cov = wb['Test Case Coverage Plan']

ws_cov.row_dimensions[1].height = 20
ws_cov.row_dimensions[2].height = 30

# Cập nhật danh sách Planned Tester lookup:
testers_cov_list = [LEADER_DINH, MEMBER_HUNG, MEMBER_NHA]
for i, pt in enumerate(testers_cov_list):
    ws_cov.cell(row=41 + i, column=1, value=pt)

for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
    cov_row = 3 + idx
    list_row = 2 + idx
    (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

    ws_cov.cell(cov_row, 1, f"='Test Case List'!C{list_row}")
    ws_cov.cell(cov_row, 2, f"='Test Case List'!D{list_row}")
    
    ws_cov.cell(cov_row, 3, iter_plan[0])
    ws_cov.cell(cov_row, 4, iter_plan[1])
    ws_cov.cell(cov_row, 5, iter_plan[2])
    ws_cov.cell(cov_row, 6, iter_plan[3])
    ws_cov.cell(cov_row, 7, iter_plan[4])

    ws_cov.cell(cov_row, 8, lead_tester)
    ws_cov.cell(cov_row, 9, "")
    ws_cov.row_dimensions[cov_row].height = 26

ws_cov.cell(28, 1, "Insert Lines Above Here Only")
ws_cov.cell(28, 2, None)
ws_cov.cell(28, 3, "Insert Lines Above Here Only")
for col in range(4, 10):
    ws_cov.cell(28, col, None)
ws_cov.row_dimensions[28].height = 22

for col in range(1, 10):
    ws_cov.cell(29, col, None)

cov_col_widths = {
    "A": 16,
    "B": 45,
    "C": 14,
    "D": 14,
    "E": 14,
    "F": 14,
    "G": 14,
    "H": 34, # Lead Tester
    "I": 22
}
for col_let, w in cov_col_widths.items():
    ws_cov.column_dimensions[col_let].width = w


# =============================================================================
# 4. CẬP NHẬT 5 SHEETS 'Test Iteration 1' ĐẾN 'Test Iteration 5'
# =============================================================================
iteration_names = [
    "Iteration One",
    "Iteration Two",
    "Iteration Three",
    "Iteration Four",
    "Iteration Five"
]

for iter_idx in range(1, 6):
    sname = f"Test Iteration {iter_idx}"
    print(f"4.{iter_idx}. Đang cập nhật Sheet '{sname}' (Tester thực tế)...")
    ws_it = wb[sname]

    ws_it['G31'] = iteration_names[iter_idx - 1]
    ws_it.row_dimensions[1].height = 30

    # Cập nhật lookup danh sách Tester ở dưới A31:A33
    for i, t_name in enumerate(testers_cov_list):
        ws_it.cell(row=31 + i, column=1, value=t_name)

    for idx, tc in enumerate(OFFICEPET_25_TEST_CASES):
        r = 2 + idx
        list_row = 2 + idx
        iter_col_idx = 2 + iter_idx
        (sc_id, sc_name, tc_id, tc_name, tc_desc, status, owner, approver, comments, ucs, frs, class_a, class_b, iter_plan, lead_tester) = tc

        in_this_iter = iter_plan[iter_idx - 1]

        ws_it.cell(r, 1, f"='Test Case List'!C{list_row}")
        ws_it.cell(r, 2, f"='Test Case List'!D{list_row}")
        ws_it.cell(r, 3, f"=VLOOKUP(A{r},'Test Case Coverage Plan'!$A$3:$G$27,{iter_col_idx},FALSE)")
        ws_it.cell(r, 4, lead_tester if in_this_iter == "Yes" else None)

        # Trạng thái: 'Not Run' nếu trong đợt test, 'N/A' nếu ngoài đợt test
        ws_it.cell(r, 5, "Not Run" if in_this_iter == "Yes" else "N/A")

        ws_it.cell(r, 6, "")
        ws_it.cell(r, 7, "")
        ws_it.cell(r, 8, "")
        ws_it.row_dimensions[r].height = 26

    ws_it.cell(27, 1, "Insert Rows Only Above Here")
    ws_it.cell(27, 2, None)
    ws_it.cell(27, 3, None)
    ws_it.cell(27, 4, "Insert Rows Only Above Here")
    for col in range(5, 9):
        ws_it.cell(27, col, None)
    ws_it.row_dimensions[27].height = 22

    for col in range(1, 9):
        ws_it.cell(28, col, None)

    it_col_widths = {
        "A": 16,
        "B": 45,
        "C": 24,
        "D": 34, # Tester
        "E": 20,
        "F": 22,
        "G": 24,
        "H": 28
    }
    for col_let, w in it_col_widths.items():
        ws_it.column_dimensions[col_let].width = w


# =============================================================================
# 5. CẬP NHẬT SHEET 'Stats'
# =============================================================================
print("5. Đang cập nhật Sheet 'Stats'...")
ws_stats = wb['Stats']
ws_stats.row_dimensions[1].height = 28

for iter_idx in range(1, 6):
    iter_sheet_name = f"Test Iteration {iter_idx}"
    start_stat_row = 2 + (iter_idx - 1) * 25

    for idx in range(25):
        stat_r = start_stat_row + idx
        source_r = 2 + idx
        list_r = 2 + idx

        ws_stats.cell(stat_r, 1, f"='{iter_sheet_name}'!A{source_r}")
        ws_stats.cell(stat_r, 2, f"='{iter_sheet_name}'!B{source_r}")
        ws_stats.cell(stat_r, 3, f"='{iter_sheet_name}'!$G$31")
        ws_stats.cell(stat_r, 4, f"='{iter_sheet_name}'!C{source_r}")
        ws_stats.cell(stat_r, 5, f"='{iter_sheet_name}'!D{source_r}")
        ws_stats.cell(stat_r, 6, f"='{iter_sheet_name}'!E{source_r}")
        ws_stats.cell(stat_r, 7, f"='Test Case List'!L{list_r}")
        ws_stats.cell(stat_r, 8, f"='Test Case List'!M{list_r}")
        ws_stats.row_dimensions[stat_r].height = 22

ws_stats.cell(127, 1, "Insert Rows Only Above Here")
ws_stats.cell(127, 5, "Insert Rows Only Above Here")
ws_stats.row_dimensions[127].height = 22

for r_clear in range(128, ws_stats.max_row + 1):
    for c_clear in range(1, 9):
        ws_stats.cell(r_clear, c_clear, None)

stat_col_widths = {
    "A": 16,
    "B": 45,
    "C": 20,
    "D": 22,
    "E": 34, # Tester
    "F": 20,
    "G": 20,
    "H": 20
}
for col_let, w in stat_col_widths.items():
    ws_stats.column_dimensions[col_let].width = w


# =============================================================================
# 6. SHEET 'Pivot Tables'
# =============================================================================
print("6. Đang căn chỉnh độ rộng cột và chiều cao hàng cho Sheet 'Pivot Tables'...")
ws_piv = wb['Pivot Tables']
for r in range(1, ws_piv.max_row + 1):
    ws_piv.row_dimensions[r].height = 22

piv_col_widths = {
    "A": 28,
    "B": 22,
    "C": 20,
    "D": 20,
    "E": 20,
    "F": 20,
    "G": 20
}
for col_let, w in piv_col_widths.items():
    ws_piv.column_dimensions[col_let].width = w


# =============================================================================
# 7. ĐỒNG BỘ 100% FONT CHỮ TIMES NEW ROMAN & BẬT HIỂN THỊ LƯỚI
# =============================================================================
print("7. Đang rà soát và ép buộc 100% font chữ Times New Roman trên toàn bộ 10 sheets...")

for sname in wb.sheetnames:
    ws = wb[sname]
    if ws.views and ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                set_tnr(cell)

                if isinstance(cell.value, str):
                    cur_align = cell.alignment
                    if '\n' in cell.value or len(cell.value) > 20:
                        cell.alignment = Alignment(
                            horizontal=cur_align.horizontal if cur_align and cur_align.horizontal else "left",
                            vertical="top" if '\n' in cell.value else (cur_align.vertical if cur_align and cur_align.vertical else "center"),
                            wrap_text=True
                        )


# =============================================================================
# 8. LƯU WORKBOOK
# =============================================================================
print(f"8. Đang lưu file chuẩn hóa vào: {FILE_PATH}...")
wb.save(FILE_PATH)
print(">>> CẬP NHẬT TEST PLAN THÀNH CÔNG ĐÚNG HOÀN TOÀN THEO HÌNH ẢNH! <<<")
